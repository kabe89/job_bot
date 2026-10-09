"""The tailored cover letter must be produced as a Word document too.

The resume has always been written as .md, .pdf and .docx, but the cover letter
was written only as {slug}_cover.txt. Employers ask for a cover letter as an
uploadable document, so a bare .txt means hand-converting it every time.

Gated on settings.output_docx, exactly like the resume's .docx.
"""
from pathlib import Path
from unittest.mock import patch

import pytest

from jobbot.pipeline import _write_cover_letter_docx


class TestWriteCoverLetterDocx:
    def test_writes_a_docx_next_to_the_txt(self, tmp_path):
        txt = tmp_path / "42_acme_scientist_cover.txt"
        txt.write_text("Dear Hiring Manager,\n\nI am applying.\n\nSincerely,\nJane",
                       encoding="utf-8")
        out = _write_cover_letter_docx("Dear Hiring Manager,\n\nI am applying.", txt)
        assert out is not None
        assert out.exists()
        assert out.name == "42_acme_scientist_cover_letter.docx"

    def test_content_survives_into_the_document(self, tmp_path):
        from docx import Document
        txt = tmp_path / "7_x_y_cover.txt"
        body = "Apex's research initiative aligns with my work.\n\nSincerely,\nJane Doe"
        out = _write_cover_letter_docx(body, txt)
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "Apex" in text
        assert "Jane Doe" in text

    def test_empty_cover_text_writes_nothing(self, tmp_path):
        txt = tmp_path / "9_a_b_cover.txt"
        assert _write_cover_letter_docx("", txt) is None
        assert _write_cover_letter_docx("   ", txt) is None

    def test_failure_is_swallowed_and_returns_none(self, tmp_path):
        """A docx render failure must never sink an otherwise good tailor run."""
        txt = tmp_path / "1_a_b_cover.txt"
        with patch("jobbot.resume.cover_letter_to_docx", side_effect=OSError("boom")):
            assert _write_cover_letter_docx("Some text", txt) is None

    def test_respects_output_docx_setting(self, tmp_path):
        txt = tmp_path / "2_a_b_cover.txt"
        with patch("jobbot.pipeline.settings") as s:
            s.output_docx = False
            assert _write_cover_letter_docx("Some text", txt) is None

    def test_formal_letterhead_structure(self, tmp_path):
        from docx import Document
        txt = tmp_path / "10_test_cover.txt"
        body = """Jane Doe, Ph.D.
123 Innovation Way
Tech City, CA 94016
October 1, 2026
Search Committee
Apex Global Technologies
San Francisco, CA 94105
Dear Search Committee,
I am excited to apply for the Senior Staff Engineer position.
Sincerely,
Jane Doe, Ph.D.
Department of Engineering & Applied Science
jane.doe@example.com | +1-555-019-2834"""
        out = _write_cover_letter_docx(body, txt)
        doc = Document(str(out))
        paragraphs = [p.text for p in doc.paragraphs]
        assert "Jane Doe, Ph.D." in paragraphs[0]
        assert "123 Innovation Way" in paragraphs[1]
        assert "October 1, 2026" in paragraphs[3]
        assert any("Senior Staff Engineer" in p for p in paragraphs)
        assert any("jane.doe@example.com" in p for p in paragraphs)


class TestCoverLetterDownloadRoute:
    def test_download_cover_docx_success(self, tmp_path):
        from jobbot import web as web_mod
        from jobbot.models import SessionLocal, Job, Application

        app = web_mod.create_app()
        app.config.update(TESTING=True)
        client = app.test_client()

        with SessionLocal() as db:
            job = Job(hash="test-hash-cover-anon", source="direct", title="Engineer", company="Acme", url="https://example.com/job/test")
            db.add(job)
            db.commit()

            txt_file = tmp_path / "app_cover.txt"
            txt_file.write_text("Cover letter text", encoding="utf-8")
            _write_cover_letter_docx("Dear Hiring Team,\n\nTest cover letter.", txt_file)

            application = Application(
                job_id=job.id,
                cover_letter_path=str(txt_file),
                status="draft",
            )
            db.add(application)
            db.commit()
            app_id = application.id

        resp = client.get(f"/download/app/{app_id}/cover-docx")
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in resp.content_type
        assert resp.headers.get("Content-Disposition", "").endswith("app_cover_letter.docx")

    def test_download_cover_docx_404_when_missing(self):
        from jobbot import web as web_mod
        app = web_mod.create_app()
        app.config.update(TESTING=True)
        client = app.test_client()

        resp = client.get("/download/app/999999/cover-docx")
        assert resp.status_code == 404
